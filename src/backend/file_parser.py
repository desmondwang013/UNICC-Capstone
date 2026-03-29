"""
File parser — converts uploaded files into plain text for evaluation.

Supported: .txt .md .json .csv .xlsx .docx
For large datasets (CSV/XLSX), samples up to MAX_BATCH_ROWS rows.
"""
from __future__ import annotations

import json
import io
from typing import Tuple
from models import ContentType

MAX_BATCH_ROWS  = 30    # max rows to sample from CSV/XLSX
MAX_CONTENT_LEN = 20000 # max characters sent to judges


def parse_file(filename: str, data: bytes) -> Tuple[str, ContentType]:
    """
    Parse uploaded file bytes into (text_content, detected_content_type).
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("txt", "md"):
        return _handle_text(data), ContentType.DOCUMENT

    if ext == "json":
        return _handle_json(data)

    if ext == "csv":
        return _handle_csv(data)

    if ext == "xlsx":
        return _handle_xlsx(data)

    if ext == "docx":
        return _handle_docx(data)

    # fallback: try raw text
    try:
        return data.decode("utf-8", errors="replace"), ContentType.DOCUMENT
    except Exception:
        raise ValueError(f"Unsupported file type: {ext}")


# ---------------------------------------------------------------------------

def _handle_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")[:MAX_CONTENT_LEN]


def _handle_json(data: bytes) -> Tuple[str, ContentType]:
    raw = json.loads(data.decode("utf-8", errors="replace"))

    # Detect conversation format: list of {role, content} dicts
    if isinstance(raw, list) and raw and isinstance(raw[0], dict):
        if "role" in raw[0] and "content" in raw[0]:
            text = _format_conversation(raw)
            return text[:MAX_CONTENT_LEN], ContentType.CONVERSATION

        # Batch of text samples: list of {id, text} or {comment_text, ...}
        text = _format_batch(raw)
        return text[:MAX_CONTENT_LEN], ContentType.BATCH

    # Single object — render as formatted text
    return json.dumps(raw, indent=2)[:MAX_CONTENT_LEN], ContentType.DOCUMENT


def _handle_csv(data: bytes) -> Tuple[str, ContentType]:
    import csv
    text_io = io.StringIO(data.decode("utf-8", errors="replace"))
    reader  = csv.DictReader(text_io)
    rows    = list(reader)

    if not rows:
        return "", ContentType.BATCH

    headers = list(rows[0].keys())

    # Detect conversation CSV: has role + content columns
    if {"role", "content"}.issubset(set(h.lower() for h in headers)):
        msgs = [{"role": r.get("role",""), "content": r.get("content","")} for r in rows]
        return _format_conversation(msgs)[:MAX_CONTENT_LEN], ContentType.CONVERSATION

    # Detect text column
    text_col = _find_text_column(headers)
    sample   = rows[:MAX_BATCH_ROWS]
    lines    = []
    for i, row in enumerate(sample):
        text_val = row.get(text_col, str(row))
        lines.append(f"[{i+1}] {text_val}")

    total = len(rows)
    header = f"=== Dataset: {total} rows total, showing {len(sample)} samples ===\n\n"
    return (header + "\n".join(lines))[:MAX_CONTENT_LEN], ContentType.BATCH


def _handle_xlsx(data: bytes) -> Tuple[str, ContentType]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return "", ContentType.BATCH

    headers = [str(h) if h is not None else "" for h in rows[0]]
    data_rows = rows[1:]

    # Detect conversation
    lower_h = [h.lower() for h in headers]
    if "role" in lower_h and "content" in lower_h:
        ri = lower_h.index("role")
        ci = lower_h.index("content")
        msgs = [{"role": str(r[ri] or ""), "content": str(r[ci] or "")} for r in data_rows]
        return _format_conversation(msgs)[:MAX_CONTENT_LEN], ContentType.CONVERSATION

    # Batch
    text_col_idx = _find_text_column_idx(headers)
    sample = data_rows[:MAX_BATCH_ROWS]
    lines  = []
    for i, row in enumerate(sample):
        val = row[text_col_idx] if text_col_idx < len(row) else str(row)
        lines.append(f"[{i+1}] {val}")

    total  = len(data_rows)
    header = f"=== Dataset: {total} rows total, showing {len(sample)} samples ===\n\n"
    return (header + "\n".join(lines))[:MAX_CONTENT_LEN], ContentType.BATCH


def _handle_docx(data: bytes) -> Tuple[str, ContentType]:
    from docx import Document
    doc   = Document(io.BytesIO(data))
    text  = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return text[:MAX_CONTENT_LEN], ContentType.DOCUMENT


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def _format_conversation(messages: list) -> str:
    """Format conversation — extract and label AI (assistant) turns."""
    lines = ["=== Conversation Log (evaluating AI/assistant responses) ===\n"]
    for msg in messages:
        role    = str(msg.get("role", "unknown")).lower()
        content = str(msg.get("content", ""))
        prefix  = "► AI:" if role == "assistant" else "  User:"
        lines.append(f"{prefix} {content}")
    return "\n".join(lines)


def _format_batch(items: list) -> str:
    """Format list of objects as numbered text samples."""
    text_key = None
    if items:
        for k in ("comment_text", "text", "content", "message", "body", "input"):
            if k in items[0]:
                text_key = k
                break

    sample = items[:MAX_BATCH_ROWS]
    lines  = [f"=== Batch: {len(items)} items total, showing {len(sample)} samples ===\n"]
    for i, item in enumerate(sample):
        if text_key:
            lines.append(f"[{i+1}] {item[text_key]}")
        else:
            lines.append(f"[{i+1}] {json.dumps(item)[:200]}")
    return "\n".join(lines)


def _find_text_column(headers: list) -> str:
    priority = ["comment_text", "text", "content", "message", "body", "input", "response"]
    lower_h  = [h.lower() for h in headers]
    for p in priority:
        if p in lower_h:
            return headers[lower_h.index(p)]
    return headers[0]


def _find_text_column_idx(headers: list) -> int:
    priority = ["comment_text", "text", "content", "message", "body", "input", "response"]
    lower_h  = [h.lower() for h in headers]
    for p in priority:
        if p in lower_h:
            return lower_h.index(p)
    return 0


def detect_content_type(text: str) -> ContentType:
    """Heuristic detection for pasted text."""
    lines = text.strip().splitlines()

    # JSON-like
    stripped = text.strip()
    if stripped.startswith("[") or stripped.startswith("{"):
        try:
            parsed = json.loads(stripped)
            if isinstance(parsed, list) and parsed:
                if "role" in parsed[0] and "content" in parsed[0]:
                    return ContentType.CONVERSATION
            return ContentType.BATCH
        except Exception:
            pass

    # Conversation markers
    conv_markers = ["user:", "assistant:", "human:", "ai:", "► ai:", "[user]", "[assistant]"]
    lower_text   = text.lower()
    if sum(1 for m in conv_markers if m in lower_text) >= 2:
        return ContentType.CONVERSATION

    # Many short lines = likely batch
    if len(lines) > 10 and all(len(l) < 300 for l in lines[:10]):
        return ContentType.BATCH

    return ContentType.DOCUMENT
